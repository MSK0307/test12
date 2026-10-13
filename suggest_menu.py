"""menu_db.xlsx の登録済みメニューを参考に、Claude に新しいメニューを提案させるツール。

使い方:
    python suggest_menu.py                          # 条件なし: DBの傾向に合わせつつ、少ない種類を補う
    python suggest_menu.py "柔らかめの魚料理"         # 条件を自由に書く
    python suggest_menu.py "薄味の副菜" -n 3 --add    # 3件提案し、メインのDBにも登録する
"""

import argparse
import sys
from datetime import datetime

import anthropic
from openpyxl.styles import Alignment, Font, PatternFill

from analyze_menu import COLUMNS, DB_PATH, SCHEMA, open_db, request_json

SUGGEST_SHEET = "提案"
SUGGEST_HEADERS = ["提案日時", "条件", "提案理由", "作り方の概要"]

PROPOSAL_SCHEMA = {
    "type": "object",
    "properties": {
        "proposals": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string", "description": "メニュー名"},
                    "reason": {"type": "string", "description": "提案理由。データベースのどの傾向や不足を踏まえたか"},
                    "recipe": {"type": "string", "description": "作り方の概要（3〜4文）"},
                    **SCHEMA["properties"],
                },
                "required": ["name", "reason", "recipe", *SCHEMA["required"]],
                "additionalProperties": False,
            },
        },
    },
    "required": ["proposals"],
    "additionalProperties": False,
}

SYSTEM_PROMPT = """あなたは献立づくりを手伝う料理の専門家です。ユーザーのメニューデータベース（食材・形態・柔らかさ・味で分析済み）を参考に、新しいメニューを提案してください。

- データベースに登録済みのメニューや、実質的に同じ料理は提案しない。
- 条件が指定されていれば、必ずそれを満たす。
- 条件がなければ、データベース全体の傾向（よく使う食材、柔らかさ、味の濃さなど）に合わせつつ、登録が少ない料理形態・食材・味を補うメニューを選ぶ。
- 提案どうしで料理形態や主な食材が偏らないようにする。
- 各提案の分析項目は、データベースと同じ基準で評価する（柔らかさは 1: とても柔らかい 〜 5: 硬い）。"""


def load_menu_table(ws) -> tuple[list[str], str]:
    """メインシートを、Claude に渡すタブ区切りの表にする。"""
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    headers = [header for header, _ in COLUMNS if header != "分析日時"]
    lines = ["\t".join(headers)]
    for row in rows:
        lines.append("\t".join("" if v is None else str(v) for v in row[:len(headers)]))
    return [row[0] for row in rows], "\n".join(lines)


def open_suggest_sheet(wb):
    if SUGGEST_SHEET in wb.sheetnames:
        return wb[SUGGEST_SHEET]
    ws = wb.create_sheet(SUGGEST_SHEET)
    ws.append(["メニュー名", *SUGGEST_HEADERS, *[h for h, _ in COLUMNS[1:-1]]])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A5C4F")
        cell.alignment = Alignment(horizontal="center")
    widths = [18, 17, 20, 40, 50, 30, 20, 12, 16, 12, 14, 18, 20, 10, 40]
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    ws.freeze_panes = "B2"
    return ws


def main() -> int:
    parser = argparse.ArgumentParser(description="menu_db.xlsx を参考に新しいメニューを提案します")
    parser.add_argument("request", nargs="?", default="", help="提案の条件（例: 柔らかめの魚料理）")
    parser.add_argument("-n", "--count", type=int, default=5, help="提案する件数（既定: 5）")
    parser.add_argument("--add", action="store_true", help="提案をメインのデータベースにも登録する")
    args = parser.parse_args()

    if not DB_PATH.exists():
        print(f"{DB_PATH.name} がありません。先に analyze_menu.py でメニューを登録してください。", file=sys.stderr)
        return 1
    wb, ws = open_db()
    existing, table = load_menu_table(ws)
    if not existing:
        print("データベースにメニューが登録されていません。先に analyze_menu.py で登録してください。", file=sys.stderr)
        return 1

    condition = args.request or "（指定なし）"
    content = (
        f"## メニューデータベース（{len(existing)}件）\n{table}\n\n"
        f"## 条件\n{condition}\n\n"
        f"新しいメニューを {args.count} 件提案してください。"
    )
    print(f"{len(existing)} 件のメニューを参考に、{args.count} 件提案します（条件: {condition}）…\n")

    client = anthropic.Anthropic()
    try:
        result = request_json(client, SYSTEM_PROMPT, content, PROPOSAL_SCHEMA,
                              effort="medium", max_tokens=16000)
    except anthropic.APIConnectionError as e:
        print(f"× API に接続できません ({e})", file=sys.stderr)
        return 1
    except anthropic.APIStatusError as e:
        print(f"× API エラー {e.status_code} ({e.message})", file=sys.stderr)
        return 1
    except (RuntimeError, ValueError) as e:
        print(f"× {e}", file=sys.stderr)
        return 1

    # 念のため、登録済みと同名の提案は除く
    proposals = [p for p in result["proposals"] if p["name"] not in existing]

    suggest_ws = open_suggest_sheet(wb)
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for i, p in enumerate(proposals, start=1):
        analysis = [get(p["name"], p) for _, get in COLUMNS]
        suggest_ws.append([p["name"], now, condition, p["reason"], p["recipe"], *analysis[1:-1]])
        if args.add:
            ws.append(analysis)
        print(f"{i}. {p['name']}  （{p['dish_form']} / 柔らかさ {p['softness']} / "
              f"{'・'.join(p['basic_tastes'])} / {p['taste_strength']}）")
        print(f"   食材: {'、'.join(p['ingredients'])}")
        print(f"   理由: {p['reason']}")
        print(f"   作り方: {p['recipe']}\n")

    if args.add:
        ws.auto_filter.ref = ws.dimensions
    wb.save(DB_PATH)
    where = "「提案」シートとメインのデータベース" if args.add else "「提案」シート"
    print(f"{len(proposals)} 件を {DB_PATH.name} の{where}に保存しました。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
