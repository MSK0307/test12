"""料理メニューを Claude で分析し、Excel (menu_db.xlsx) に追記するツール。

使い方:
    python analyze_menu.py                       # 入力フォーム（料理名・材料・調味料を順に入力）
    python analyze_menu.py -f menus.csv          # CSV で一括登録（列: 料理名,材料,調味料）
    python analyze_menu.py -f menus.csv --force  # 登録済みのメニューも再分析する
    python analyze_menu.py 肉じゃが 鯖の味噌煮     # 料理名だけで手早く登録
"""

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import anthropic
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MODEL = "claude-opus-5-5"
DB_PATH = Path(__file__).with_name("menu_db.xlsx")
SHEET_NAME = "メニュー分析"

INGREDIENT_CATEGORIES = [
    "肉", "魚介", "野菜", "いも", "きのこ", "海藻", "豆・大豆製品",
    "卵", "乳製品", "穀物", "果物", "その他",
]
DISH_FORMS = [
    "汁物", "煮物", "焼き物", "炒め物", "揚げ物", "蒸し物", "和え物・サラダ",
    "生もの・刺身", "鍋物", "ご飯もの", "麺類", "パン・粉もの", "デザート", "その他",
]
SHAPES = ["塊・大きめ", "一口大", "薄切り・細切り", "刻み・ほぐし", "ペースト・ムース状", "液状・汁気多め"]
SOFTNESS_LABELS = {1: "とても柔らかい", 2: "柔らかい", 3: "普通", 4: "やや硬い", 5: "硬い"}
BASIC_TASTES = ["甘味", "塩味", "酸味", "苦味", "うま味", "辛味"]
STRENGTHS = ["薄味", "普通", "濃いめ"]

# Excel の列（見出し, 分析結果から値を取り出す関数）
INPUT_COLUMNS = [
    ("入力: 材料", lambda name, r: r.get("input_ingredients", "")),
    ("入力: 調味料", lambda name, r: r.get("input_seasonings", "")),
]
ANALYSIS_COLUMNS = [
    ("主な食材", lambda name, r: "、".join(r["ingredients"])),
    ("食材カテゴリ", lambda name, r: "、".join(r["ingredient_categories"])),
    ("料理形態", lambda name, r: r["dish_form"]),
    ("形状", lambda name, r: r["shape"]),
    ("柔らかさ(1-5)", lambda name, r: r["softness"]),
    ("柔らかさ", lambda name, r: SOFTNESS_LABELS[r["softness"]]),
    ("基本味", lambda name, r: "、".join(r["basic_tastes"])),
    ("味付け", lambda name, r: "、".join(r["seasonings"])),
    ("味の濃さ", lambda name, r: r["taste_strength"]),
    ("補足", lambda name, r: r["notes"]),
]
COLUMNS = [
    ("メニュー名", lambda name, r: name),
    *INPUT_COLUMNS,
    *ANALYSIS_COLUMNS,
    ("分析日時", lambda name, r: datetime.now().strftime("%Y-%m-%d %H:%M")),
]

SCHEMA = {
    "type": "object",
    "properties": {
        "ingredients": {
            "type": "array", "items": {"type": "string"},
            "description": "一般的なレシピで使われる主な食材（調味料は除く）",
        },
        "ingredient_categories": {
            "type": "array", "items": {"type": "string", "enum": INGREDIENT_CATEGORIES},
        },
        "dish_form": {"type": "string", "enum": DISH_FORMS},
        "shape": {"type": "string", "enum": SHAPES},
        "softness": {"type": "integer", "enum": list(SOFTNESS_LABELS)},
        "basic_tastes": {
            "type": "array", "items": {"type": "string", "enum": BASIC_TASTES},
            "description": "はっきり感じられる基本味。強いものから順に",
        },
        "seasonings": {
            "type": "array", "items": {"type": "string"},
            "description": "主な味付け（醤油、味噌、トマトソースなど）",
        },
        "taste_strength": {"type": "string", "enum": STRENGTHS},
        "notes": {"type": "string", "description": "判定の前提や、作り方で変わる点を一言で"},
    },
    "required": [
        "ingredients", "ingredient_categories", "dish_form", "shape", "softness",
        "basic_tastes", "seasonings", "taste_strength", "notes",
    ],
    "additionalProperties": False,
}

SYSTEM_PROMPT = """あなたは料理の分析担当です。メニュー名から、日本の家庭や飲食店で一般的に作られる形を想定して、その料理を分析してください。

- 食材: 主な食材を挙げ、該当する食材カテゴリをすべて選ぶ。
- 形態: 調理法から料理形態を1つ、料理に入っている主な具の形状を1つ選ぶ。
- 柔らかさ: 料理全体の食感を 1(とても柔らかい: 舌でつぶせる) 〜 5(硬い: よく噛む必要がある) で評価する。
- 味: はっきり感じられる基本味、主な味付け、味の濃さを答える。
- メニュー名があいまいな場合は、最も一般的な解釈で答え、その前提を notes に書く。
- 材料や調味料が指定されている場合は、それを使って作った料理として分析する。主な食材・味付けには指定されたものを必ず含め、一般的なレシピとの違いがあれば notes に書く。"""


def request_json(client: anthropic.Anthropic, system: str, content: str, schema: dict,
                 effort: str, max_tokens: int = 8000) -> dict:
    """Claude に schema どおりの JSON で答えさせ、dict にして返す。"""
    response = client.beta.messages.create(
        model=MODEL,
        max_tokens=max_tokens,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        output_config={
            "effort": effort,
            "format": {"type": "json_schema", "schema": schema},
        },
        system=system,
        messages=[{"role": "user", "content": content}],
    )
    if response.stop_reason == "refusal":
        raise RuntimeError("Claude が回答を断りました")
    if response.stop_reason == "max_tokens":
        raise RuntimeError("出力が上限に達して途中で切れました")
    text = next(b.text for b in response.content if b.type == "text")
    return json.loads(text)


def analyze(client: anthropic.Anthropic, menu_name: str, ingredients: str = "",
            seasonings: str = "") -> dict:
    content = f"メニュー名: {menu_name}"
    if ingredients:
        content += f"\n材料: {ingredients}"
    if seasonings:
        content += f"\n調味料: {seasonings}"
    return request_json(client, SYSTEM_PROMPT, content, SCHEMA, effort="low")


def normalize_list(text: str) -> str:
    """「じゃがいも, 牛肉 玉ねぎ」のような入力を「じゃがいも、牛肉、玉ねぎ」にそろえる。"""
    return "、".join(item for item in re.split(r"[、,，・\s]+", text or "") if item)


def open_db() -> tuple[Workbook, object]:
    if DB_PATH.exists():
        wb = load_workbook(DB_PATH)
        ws = wb[SHEET_NAME]
        headers = [cell.value for cell in ws[1]]
        if headers[:len(COLUMNS)] != [header for header, _ in COLUMNS]:
            sys.exit(f"{DB_PATH.name} の列の構成が現在のツールと違います。"
                     "ファイル名を変えて退避してから、もう一度実行してください。")
        return wb, ws
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append([header for header, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6D7A")
        cell.alignment = Alignment(horizontal="center")
    widths = [18, 24, 20, 30, 20, 12, 16, 12, 14, 18, 20, 10, 40, 17]
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return wb, ws


def find_row(ws, menu_name: str) -> int | None:
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == menu_name:
            return row
    return None


def read_csv(path: Path) -> list[dict]:
    """CSV（列: 料理名,材料,調味料）を読む。Excel で保存した Shift_JIS の CSV にも対応する。"""
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp932")
    entries = []
    for row in csv.reader(text.splitlines()):
        cells = [cell.strip() for cell in row] + ["", "", ""]
        if not cells[0] or cells[0] in ("料理名", "メニュー名"):  # 空行と見出し行は飛ばす
            continue
        entries.append({"name": cells[0], "ingredients": cells[1], "seasonings": cells[2]})
    return entries


def ask(prompt: str) -> str:
    try:
        return input(prompt).strip()
    except EOFError:
        return ""


def ask_entry() -> dict | None:
    """入力フォーム。料理名が空のままなら聞き直し、q で終了する。"""
    print("\n--- 新しいメニュー（終了するときは料理名に q と入力）---")
    while True:
        name = ask("料理名（必須）: ")
        if name.lower() == "q":
            return None
        if name:
            break
        print("  料理名は必須です。")
    ingredients = ask("材料（任意・Enterで省略）: ")
    seasonings = ask("調味料（任意・Enterで省略）: ")
    return {"name": name, "ingredients": ingredients, "seasonings": seasonings}


def register(client: anthropic.Anthropic, wb: Workbook, ws, entry: dict, overwrite: bool) -> str:
    """1件を分析して Excel に追記（overwrite なら上書き）する。結果を saved / failed / auth で返す。"""
    name = entry["name"]
    ingredients = normalize_list(entry.get("ingredients", ""))
    seasonings = normalize_list(entry.get("seasonings", ""))
    print(f"  {name} を分析しています…")
    try:
        result = analyze(client, name, ingredients, seasonings)
    except anthropic.APIConnectionError as e:
        print(f"× {name}: API に接続できません ({e})", file=sys.stderr)
        return "failed"
    except anthropic.APIStatusError as e:
        print(f"× {name}: API エラー {e.status_code} ({e.message})", file=sys.stderr)
        return "auth" if isinstance(e, anthropic.AuthenticationError) else "failed"
    except (RuntimeError, json.JSONDecodeError) as e:
        print(f"× {name}: {e}", file=sys.stderr)
        return "failed"

    result |= {"input_ingredients": ingredients, "input_seasonings": seasonings}
    values = [get(name, result) for _, get in COLUMNS]
    existing = find_row(ws, name)
    if existing and overwrite:
        for col, value in enumerate(values, start=1):
            ws.cell(row=existing, column=col, value=value)
    else:
        ws.append(values)
    ws.auto_filter.ref = ws.dimensions
    wb.save(DB_PATH)  # 1件ごとに保存し、途中で止まっても結果が残るようにする
    print(f"○ {name}: {result['dish_form']} / 柔らかさ {result['softness']}（{SOFTNESS_LABELS[result['softness']]}）/ "
          f"{'・'.join(result['basic_tastes'])} / {result['taste_strength']}")
    print(f"  食材: {'、'.join(result['ingredients'])}  味付け: {'、'.join(result['seasonings'])}")
    return "saved"


def run_form(client: anthropic.Anthropic, wb: Workbook, ws) -> int:
    saved = 0
    while (entry := ask_entry()) is not None:
        overwrite = False
        if find_row(ws, entry["name"]):
            if ask(f"  「{entry['name']}」は登録済みです。上書きしますか？ [y/N]: ").lower() != "y":
                continue
            overwrite = True
        status = register(client, wb, ws, entry, overwrite)
        if status == "auth":
            return 1
        saved += status == "saved"
    print(f"\n終了: {saved} 件を {DB_PATH.name} に保存しました。")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="料理メニューを分析して menu_db.xlsx に保存します")
    parser.add_argument("menus", nargs="*", help="メニュー名（複数可）。省略すると入力フォームが開きます")
    parser.add_argument("-f", "--file", type=Path, help="CSV ファイル（列: 料理名,材料,調味料。材料・調味料は空欄可）")
    parser.add_argument("--force", action="store_true", help="登録済みのメニューも再分析して上書きする")
    args = parser.parse_args()

    entries = [{"name": m} for m in args.menus if m.strip()]
    if args.file:
        entries += read_csv(args.file)

    client = anthropic.Anthropic()
    wb, ws = open_db()
    if not entries:
        return run_form(client, wb, ws)

    saved = failed = skipped = 0
    seen = set()
    for entry in entries:
        name = entry["name"]
        if name in seen:
            continue
        seen.add(name)
        if find_row(ws, name) and not args.force:
            print(f"- {name}: 登録済みのためスキップ（再分析は --force）")
            skipped += 1
            continue
        status = register(client, wb, ws, entry, overwrite=args.force)
        if status == "auth":
            failed += 1
            break  # 認証エラーは以降も必ず失敗するので打ち切る
        saved += status == "saved"
        failed += status == "failed"

    print(f"\n完了: 保存 {saved} 件、スキップ {skipped} 件、失敗 {failed} 件 → {DB_PATH.name}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
